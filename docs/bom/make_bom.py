"""Build the Navyam competition-drone BOM workbook."""
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
CHECKED = "31 Oct 2026"

FONT = "Arial"
H_FILL = PatternFill("solid", fgColor="1F2A33")
H_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
INPUT = Font(name=FONT, size=10, color="0000FF")
LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
TITLE = Font(name=FONT, size=14, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color="555555")
GREEN = PatternFill("solid", fgColor="DFF3E8")
AMBER = PatternFill("solid", fgColor="FDF0D5")
RED = PatternFill("solid", fgColor="FDE4E1")
GREY = PatternFill("solid", fgColor="EEF1F3")
THIN = Side(style="thin", color="C9D1D8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
INR = '₹#,##0;(₹#,##0);"-"'
INR2 = '₹#,##0.00;(₹#,##0.00);"-"'


def header(ws, row, cols, widths):
    for i, (name, width) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill, c.alignment, c.border = H_FONT, H_FILL, Alignment(wrap_text=True, vertical="center"), BOX
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30


def body_row(ws, row, values, fill=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.alignment, c.border = BODY, WRAP, BOX
        if fill:
            c.fill = fill


wb = Workbook()

# ---------------------------------------------------------------- Buy List
ws = wb.active
ws.title = "Buy List"
ws["A1"] = "Navyam competition drone: parts to buy"
ws["A1"].font = TITLE
ws["A2"] = (f"Stock and prices checked live on {CHECKED} (prices incl. GST). Blue numbers are inputs you can edit; "
            "totals recalculate. Green = buy now, amber = buy only after the check in the 'When to buy' column.")
ws["A2"].font = NOTE
ws.merge_cells("A2:L2")

cols = ["#", "Item", "What it does on the drone", "Qty", "Unit price (₹)", "Line total (₹)", "Shop",
        "Link", f"Stock on {CHECKED}", "Key specs (verified)", "Why this part", "When to buy"]
widths = [4, 30, 30, 6, 13, 13, 9, 40, 14, 44, 44, 30]
header(ws, 4, cols, widths)

items = [
    ("SimpliFly MK7 Pro 7\" carbon frame", "The airframe: arms, plates, stack mount", 1, 5719, "Robu",
     "https://robu.in/product/simplifly-mk7-pro-7-inch-drone-frame-standard/", "In stock",
     "Carbon fibre, 5 mm arms, 3 mm plates, 225 g. SimpliFly datasheet (MK7): 335 mm motor-to-motor, "
     "30.5×30.5 mm stack, recommended 2806.5 1350 KV motors, 40–50 A ESC, 7\" props. Includes 2 battery straps.",
     "Stiff carbon (no arm flex = less vibration). Built for exactly our motor/ESC/prop class. "
     "Fits the 61 cm launch box: ≤ 335 + 178 mm props + guards ≈ 54 cm.", "Now", "now"),
    ("Emax ECO II 2807 1300 KV motor", "Propulsion (4 + 1 spare)", 5, 2259, "Robu",
     "https://robu.in/product/emax-ecoii-2807-1300kv-brushless-motor/", "In stock",
     "6S, 19×19 mm M3 mount, M5 shaft, 47.6 g, max thrust 2,640 g (retailer figure), 12 mm bearing.",
     "Same 28 mm class the frame is designed for; 19×19 mount like other 2806.5/2807 motors. "
     "About 5:1 thrust-to-weight at ~2.0 kg.", "Now", "now"),
    ("T-Motor V50A 6S 4-in-1 ESC (32-bit)", "Drives the 4 motors", 1, 8479, "Robu",
     "https://robu.in/product/t-motor-v50a-6s-4in1-32bit-esc-50a/", "In stock",
     "50 A continuous / 60 A burst per motor, 3–6S, built-in current sensor, 30.5×30.5 mm M3, 18.5 g.",
     "Rated above the motor's full-throttle draw, so no software current limit is needed. Matches the frame's 30.5 stack.",
     "Now", "now"),
    ("Benewake TF-Luna LiDAR rangefinder", "Measures height above the floor (downward)", 1, 2589, "Robu",
     "https://robu.in/product/benewake-tf-luna-lidar-distance-sensor/", "In stock",
     "0.2–8 m, 1 cm resolution, UART, ≤0.35 W, up to 250 Hz.",
     "Stable altitude hold indoors. Position comes from the 2D LiDAR SLAM, so optical flow is not needed.",
     "Now", "now"),
    ("Jumper AION RX Mini ELRS receiver", "Safety pilot link + failsafe (TX16S)", 1, 1999, "Robu",
     "https://robu.in/product/jumper-rc-elrs-aion-rx-mini-receiver/", "In stock",
     "2.4 GHz ExpressLRS, 16 channels, 3.5–6 V, 1.1 g.",
     "Proper signal-loss failsafe and enough channels for kill, mode and abort switches.",
     "After checking the TX16S ELRS version is 3.x", "check"),
    ("XH-M404 (XL4016) 8 A buck converter", "Powers the Jetson from 6S (set to 19.0 V)", 2, 239, "Robu",
     "https://robu.in/product/xl4016e1-200w-step-down-power-supply-module/", "In stock",
     "4–40 V in, 1.25–36 V adjustable out, 8 A max, 200 W max, 94% efficiency.",
     "Jetson input range is 9–20 V; it draws ~2 A, a quarter of this rating. Second unit is a spare.",
     "Now", "now"),
    ("Pro-Range PETG-HS 1.75 mm filament, 1 kg, black", "Printing the deck, mounts, legs and prop guards", 1, 780, "Robu",
     "https://robu.in/product/pro-range-petg-hs-175mm-3d-printing-filament-1kg-spool-midnight-black/", "In stock",
     "PETG, 1.75 mm, 1 kg spool.",
     "PETG does not soften near warm motors like PLA does. The printed set is ~250 g, so 1 kg covers reprints.",
     "Now (skip if the lab supplies PETG)", "now"),
    ("7\" tri-blade props for 6S, e.g. HQProp 7x4x3 or 7x3.5x3 (2CW+2CCW)", "Propellers (3 sets = 1 + 2 spare)", 3, 370, "Zbotic",
     "https://zbotic.in/?s=7+inch+propeller&post_type=product", "OUT OF STOCK at Robu, Zbotic and Probots",
     "Must be 7\" (frame spec), 5 mm hole, rated for 6S 2807-class motors. Last Zbotic price for HQProp 7x4x3: ₹369.60/set.",
     "The frame maker specifies 7\". 7.5\" (HQProp 7.5x3.7x3, ₹577.50, in stock at Zbotic) is only OK if the measured "
     "gap between neighbouring motors is ≥ 207 mm (prop + 2 × guard clearance).",
     "After measuring the frame; buy whichever size the measurement allows", "check"),
]

first = 5
for n, it in enumerate(items):
    r = first + n
    name, purpose, qty, price, shop, url, stock, specs, why, when, kind = it
    fill = GREEN if kind == "now" else AMBER
    body_row(ws, r, [n + 1, name, purpose, qty, price, f"=D{r}*E{r}", shop, url, stock, specs, why, when], fill)
    ws.cell(row=r, column=4).font = INPUT
    ws.cell(row=r, column=5).font = INPUT
    ws.cell(row=r, column=5).number_format = INR2
    ws.cell(row=r, column=6).number_format = INR2
    link = ws.cell(row=r, column=8)
    link.hyperlink, link.font = url, LINK
    if "OUT OF STOCK" in stock:
        ws.cell(row=r, column=9).fill = RED
    ws.row_dimensions[r].height = 92
last = first + len(items) - 1
ws.cell(row=last, column=5).comment = Comment(
    "Last listed price at Zbotic (HQProp 7x4x3, out of stock on 31 Oct 2026). Update when you buy.", "Navyam")

t = last + 2
labels = [
    ("Buy now (green rows)", f'=SUMIFS(F{first}:F{last},L{first}:L{last},"Now*")'),
    ("Buy after a check (amber rows)", f'=SUM(F{first}:F{last})-F{t}'),
    ("Total parts", f"=F{t}+F{t + 1}"),
]
for i, (label, formula) in enumerate(labels):
    r = t + i
    ws.cell(row=r, column=5, value=label).font = BOLD
    c = ws.cell(row=r, column=6, value=formula)
    c.font, c.number_format, c.border = BOLD, INR2, BOX
ws.cell(row=t + 3, column=2,
        value="Shipping: Robu delivers free above ₹999. Check Zbotic's shipping at checkout.").font = NOTE
ws.freeze_panes = "C5"

# ---------------------------------------------------------------- Already Owned
wo = wb.create_sheet("Already Owned")
wo["A1"] = "Parts we already own that move to the competition drone (₹0 to buy)"
wo["A1"].font = TITLE
header(wo, 3, ["Item", "From BOM / team", "Use on the new drone", "Note"], [36, 22, 40, 50])
owned = [
    ("Pixhawk Cube Orange + carrier", "BOM line 1", "Flight controller", "Confirm board name in Mission Planner before flashing (Cube Orange vs Orange+)."),
    ("NVIDIA Jetson Orin Nano Super Dev Kit (8 GB)", "BOM line 5", "Onboard computer: SLAM, navigation, survivor detection", "Input 9–20 V DC jack; we feed 19 V from the buck converter."),
    ("Intel RealSense D435i", "BOM line 15", "Survivor detection + low-obstacle depth", "Front mount tilted 20–25° down (printed)."),
    ("LDROBOT D500 LiDAR", "BOM line 48", "2D SLAM: position and the map", "On a printed mast above the props."),
    ("GenX 6S 5200 mAh 40C LiPo ×2", "BOM line 3", "Flight battery", "~650–726 g. Strapped with the frame's included straps."),
    ("APM/Pixhawk power module, 90 A, 28 V", "BOM line 36", "Battery voltage/current to the Cube", "Rated for 6S."),
    ("3DR mini telemetry radio", "BOM line 26", "MAVLink to Mission Planner", ""),
    ("RadioMaster TX16S MKII ELRS", "BOM line 38", "Safety pilot transmitter", "Pairs with the new ELRS receiver."),
    ("Samsung T7 1 TB SSD", "BOM line 19", "Jetson boot drive for now", "NVMe drives were out of stock on Robu."),
    ("DC barrel pigtail for the Jetson", "On current drone", "Buck converter → Jetson DC jack", "Reuse the one already on the drone."),
    ("XT60 connectors, silicone wire, M3 screws/nuts, soldering kit", "BOM lines 7–10, 31–35, 39–41", "Wiring and assembly", ""),
]
for i, row in enumerate(owned):
    body_row(wo, 4 + i, list(row))

# ---------------------------------------------------------------- Compatibility
wc = wb.create_sheet("Compatibility Checks")
wc["A1"] = "How each part was matched to the others"
wc["A1"].font = TITLE
header(wc, 3, ["Check", "Requirement", "Our part", "Result", "Source"], [26, 36, 44, 12, 46])
checks = [
    ("Launch box", "Whole drone inside 610 × 610 mm", "Max width ≤ motor-to-motor 335 mm + prop 178 mm + guards ≈ 540 mm", "PASS", "NIDAR brief §9; SimpliFly MK7 datasheet"),
    ("Corridor", "Safe in 1 m wide corridors", "≈ 23 cm clearance each side", "PASS", "NIDAR brief §3"),
    ("Motor class", "Frame designed for 2806.5 1350 KV", "Emax 2807 1300 KV (same 28 mm class)", "PASS", "SimpliFly MK7 datasheet"),
    ("Motor mount", "Frame arm holes for 28xx motors", "19×19 mm M3", "CHECK", "Datasheet drawing has no hole spacing: fit one motor on arrival"),
    ("ESC current", "≥ motor draw; HGLRC/Emax class wants > 45 A", "V50A: 50 A cont / 60 A burst", "PASS", "Robu V50A listing"),
    ("ESC mount", "30.5 × 30.5 mm stack", "V50A: 30.5 × 30.5 mm M3", "PASS", "SimpliFly datasheet; T-Motor V50A listings"),
    ("Battery voltage", "6S on motors, ESC, power module", "Motors 6S, ESC 3–6S, power module ≤ 28 V", "PASS", "Listings"),
    ("Props", "Frame spec 7\"; 5 mm hole; 6S-rated", "7\" tri-blade (to buy when in stock)", "CHECK", "Out of stock on 31 Oct 2026"),
    ("Jetson power", "9–20 V DC jack, ~25 W", "Buck set to 19.0 V, 8 A max", "PASS", "NVIDIA carrier board spec / forum"),
    ("Thrust", "Thrust-to-weight ≥ 2", "≈ 2.0 kg all-up vs 4 × ~2.6 kg max", "PASS", "Retailer thrust figure; weight estimate (weigh after build)"),
    ("RC link", "ELRS major version must match", "TX16S module ↔ AION RX Mini", "CHECK", "Check TX16S ELRS version before binding"),
    ("Prop guards", "Mandatory full guards", "3D-printed PETG (designed in repo hardware/)", "TO DO", "NIDAR brief §2"),
]
for i, row in enumerate(checks):
    r = 4 + i
    body_row(wc, r, list(row))
    res = wc.cell(row=r, column=4)
    res.font = BOLD
    res.fill = GREEN if row[3] == "PASS" else AMBER

# ---------------------------------------------------------------- Checklist
wk = wb.create_sheet("Order & Arrival Checklist")
wk["A1"] = "Do these in order"
wk["A1"].font = TITLE
header(wk, 3, ["Step", "When", "What to do", "Done (✓)"], [6, 18, 80, 10])
steps = [
    ("Before ordering", "Check the TX16S ELRS firmware version (Model setup → ExpressLRS Lua script). It must be 3.x to match the AION RX Mini."),
    ("Order", "Order all green rows on the Buy List from Robu (one order, free shipping above ₹999)."),
    ("Frame arrives", "Fit one Emax motor on one arm with M3×8 screws. Confirm 19×19 holes line up. Do NOT open the other motors until this passes."),
    ("Frame arrives", "Fill in the 'Frame Measurements' sheet with a ruler or calipers. Send it to Claude to finalise the printed parts."),
    ("Frame arrives", "From the measured gap between neighbouring motors, decide the prop size: 7\" always works; 7.5\" only if the gap ≥ 207 mm."),
    ("Props", "Order 3 sets of the chosen props (7\" tri-blade 6S, 5 mm hole)."),
    ("Print", "Print the parts in PETG: 3 walls, 30% infill. Weigh the full printed set (target ≤ 250 g)."),
    ("Bench", "Set the buck converter to 19.0 V with a multimeter BEFORE connecting the Jetson. Run the Jetson in MAXN mode for 30 min and check the converter stays cool."),
    ("Bench", "Props OFF: ESC/motor direction test, RC failsafe test, battery failsafe test in Mission Planner."),
    ("First flight", "Weigh the complete drone with battery and update the thrust check."),
]
for i, (when, what) in enumerate(steps):
    body_row(wk, 4 + i, [i + 1, when, what, ""])

# ---------------------------------------------------------------- Measurements
wm = wb.create_sheet("Frame Measurements")
wm["A1"] = "Measure the MK7 Pro when it arrives (millimetres)"
wm["A1"].font = TITLE
wm["A2"] = "Fill the yellow cells. Example values shown in the last column are what we assumed in the design."
wm["A2"].font = NOTE
header(wm, 3, ["#", "Measurement", "How to measure", "Your value (mm)", "Design assumption (mm)"], [4, 34, 60, 16, 20])
meas = [
    ("Front-left to front-right motor centre", "Centre of motor shaft hole to centre of the other, across the front", 237),
    ("Front-left to rear-left motor centre", "Same, along the left side", 237),
    ("Diagonal motor-to-motor", "Front-left to rear-right (datasheet says 335)", 335),
    ("Arm width at motor pad", "Across the arm where the motor sits", 30),
    ("Motor hole spacing on arm", "Between two opposite M3 holes on one arm pad", 19),
    ("Top plate length", "Front to back of the top carbon plate", 150),
    ("Top plate width", "Left to right", 40),
    ("Top plate M3 hole positions", "Distance between the 4 outermost holes (length × width)", "120 × 30"),
    ("Height: bottom plate to top plate", "Outside to outside", 43),
]
for i, (m, how, assumed) in enumerate(meas):
    r = 4 + i
    body_row(wm, r, [i + 1, m, how, None, assumed])
    wm.cell(row=r, column=4).fill = PatternFill("solid", fgColor="FFFF00")
    wm.cell(row=r, column=4).font = INPUT

# ---------------------------------------------------------------- Not buying
wn = wb.create_sheet("Not Buying (Why)")
wn["A1"] = "Options considered and rejected"
wn["A1"].font = TITLE
header(wn, 3, ["Item", "Price seen (₹)", "Why not"], [44, 14, 80])
rejected = [
    ("Tarot TL65B01 650 mm frame (team proposal)", 9799, "650 mm motor spacing + props ≈ 98 cm wide: does not fit the 61 cm launch box."),
    ("5 × Hobbywing XRotor 40A ESC (team proposal)", 13050, "We already own 4 Holybro 40 A ESCs + a 4-in-1; and the new build uses one 4-in-1."),
    ("Pro-Range 4S 5200 mAh battery (team proposal)", 5779, "Our motors are 6S; 4S would give ~45% of the thrust. We own two 6S 5200 packs."),
    ("SimpliFly F450 / S500 frames", 923, "Nylon arms flex (vibration) and motor holes don't match 19×19 motors without drilling."),
    ("JHEMCU 40 A 4-in-1 (owned) on the new motors", 0, "Below the 2807's full-throttle draw; would need a software current cap."),
    ("MicoAir MTF-01 / Holybro H-Flow optical flow", 17999, "Not in stock / very expensive; LiDAR SLAM + TF-Luna cover position and height."),
    ("NVMe SSD for the Jetson", 3239, "Out of stock on Robu; boot from the owned Samsung T7 for now."),
]
for i, (item, price, why) in enumerate(rejected):
    r = 4 + i
    body_row(wn, r, [item, price, why])
    wn.cell(row=r, column=2).number_format = INR

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False

wb.calculation.fullCalcOnLoad = True
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", OUT)
